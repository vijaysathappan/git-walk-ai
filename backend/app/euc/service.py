"""Versioned EUC ingestion, analysis orchestration, and normalized persistence."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import time
import uuid
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any

from openpyxl import Workbook, load_workbook

from .. import database
from ..config import settings
from ..observability import record_audit_event, record_metric
from ..services.semantic_ledger_service import ledger_for_connection
from .analyzers import FormulaAnalyzer, OpenXMLAnalyzer, WorkbookAnalyzer
from .fingerprint import content_fingerprint, structure_fingerprint
from .models import AnalysisContext, AnalysisWarning
from .validation import EUCValidationError, validate_euc_file
from ..access_control.engine import authorization_engine, repository_resource


def _id(prefix: str, size: int = 18) -> str:
    return f"{prefix}_{uuid.uuid4().hex[:size].upper()}"


def _repository_access(conn, repository_id: str, user_id: str, edit: bool = False):
    row = conn.execute(
        """
        SELECT R.*, COALESCE(M.ROLE, CASE WHEN R.CREATED_BY=? THEN 'owner' END) AS ACCESS_ROLE
        FROM WORKBOOK_REPOSITORIES R
        LEFT JOIN REPOSITORY_MEMBERS M ON M.REPOSITORY_ID=R.REPOSITORY_ID AND M.USER_ID=?
        WHERE R.REPOSITORY_ID=? AND R.STATUS='ACTIVE'
        """,
        (user_id, user_id, repository_id),
    ).fetchone()
    if not row:
        raise PermissionError("You do not have access to this repository.")
    authorization_engine.require(
        user_id, "euc.analyze" if edit else "euc.read",
        repository_resource(repository_id, organization_id=row["ORGANIZATION_ID"]), conn=conn,
    )
    return row


def _source_commit(conn, repository_id: str) -> str | None:
    row = conn.execute(
        """SELECT B.HEAD_COMMIT_ID FROM WORKBOOK_REPOSITORIES R
           JOIN BRANCHES B ON B.BRANCH_ID=R.DEFAULT_BRANCH_ID
           WHERE R.REPOSITORY_ID=?""",
        (repository_id,),
    ).fetchone()
    return row[0] if row else None


def ingest_euc(repository_id: str, filename: str, payload: bytes, user_id: str) -> dict[str, Any]:
    validation = validate_euc_file(filename, payload)
    file_hash = content_fingerprint(payload)
    structure_hash = structure_fingerprint(payload, validation.file_type)
    conn = database._get_connection()
    try:
        repository = _repository_access(conn, repository_id, user_id, edit=True)
        existing = conn.execute(
            "SELECT * FROM EUC_ASSETS WHERE REPOSITORY_ID=? AND FILE_HASH=?",
            (repository_id, file_hash),
        ).fetchone()
        if existing:
            return _asset_dict(existing, deduplicated=True)
        ledger = ledger_for_connection(conn)
        stored = ledger.objects.put_bytes(conn, "EUC_SOURCE_FILE", payload)
        now = database._utcnow()
        euc_id = _id("EUC")
        conn.execute(
            """
            INSERT INTO EUC_ASSETS
                (EUC_ID,REPOSITORY_ID,ORIGINAL_FILENAME,FILE_TYPE,MIME_TYPE,
                 ORIGINAL_OBJECT_HASH,FILE_HASH,STRUCTURE_HASH,SIZE_BYTES,
                 ANALYSIS_STATUS,CREATED_BY,CREATED_AT,UPDATED_AT)
            VALUES (?,?,?,?,?,?,?,?,?,'PENDING',?,?,?)
            """,
            (euc_id, repository_id, Path(filename).name, validation.file_type,
             validation.mime_type, stored["object_hash"], file_hash, structure_hash,
             len(payload), user_id, now, now),
        )
        conn.execute(
            """INSERT INTO OBJECT_REFERENCES
               (PARENT_HASH,CHILD_HASH,REFERENCE_TYPE,POSITION) VALUES (?,?,?,0)
               ON CONFLICT DO NOTHING""",
            (f"EUC:{euc_id}", stored["object_hash"], "EUC_SOURCE"),
        )
        conn.commit()
        record_audit_event("EUC_IMPORTED", actor_user_id=user_id,
                           repository_id=repository_id,
                           payload={"euc_id": euc_id, "filename": Path(filename).name,
                                    "file_hash": file_hash, "deduplicated": not stored["created"]})
        return {"euc_id": euc_id, "repository_id": repository_id,
                "filename": Path(filename).name, "file_type": validation.file_type,
                "file_hash": file_hash, "structure_hash": structure_hash,
                "size_bytes": len(payload), "status": "PENDING",
                "has_macros": validation.has_macros,
                "deduplicated": not stored["created"], "repository_name": repository["REPOSITORY_NAME"]}
    finally:
        conn.close()


def _asset_dict(row, deduplicated: bool = False) -> dict[str, Any]:
    return {"euc_id": row["EUC_ID"], "repository_id": row["REPOSITORY_ID"],
            "filename": row["ORIGINAL_FILENAME"], "file_type": row["FILE_TYPE"],
            "file_hash": row["FILE_HASH"], "structure_hash": row["STRUCTURE_HASH"],
            "size_bytes": row["SIZE_BYTES"], "status": row["ANALYSIS_STATUS"],
            "latest_analysis_id": row["LATEST_ANALYSIS_ID"], "created_at": row["CREATED_AT"],
            "deduplicated": deduplicated}


def _stable_sheets(conn, repository_id: str) -> dict[str, str]:
    return {row["SHEET_NAME"]: row["SHEET_ID"] for row in conn.execute(
        "SELECT SHEET_ID,SHEET_NAME FROM WORKBOOK_SHEETS WHERE REPOSITORY_ID=? AND STATUS='ACTIVE'",
        (repository_id,),
    )}


def _workbook_for(path: Path, file_type: str):
    if file_type != "csv":
        return load_workbook(path, data_only=False, keep_links=True, keep_vba=file_type == "xlsm")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = path.stem[:31] or "CSV"
    cell_count = 0
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        for row in csv.reader(handle):
            cell_count += len(row)
            if cell_count > settings.euc_max_cells:
                raise EUCValidationError("OVERSIZED_WORKBOOK", "CSV cell limit exceeded.")
            sheet.append(row)
    return workbook


def analyze_euc(euc_id: str, user_id: str) -> dict[str, Any]:
    conn = database._get_connection()
    analysis_id = _id("ANA")
    started = time.perf_counter()
    temp_path = None
    try:
        asset = conn.execute("SELECT * FROM EUC_ASSETS WHERE EUC_ID=?", (euc_id,)).fetchone()
        if not asset:
            raise KeyError("EUC asset does not exist.")
        _repository_access(conn, asset["REPOSITORY_ID"], user_id, edit=True)
        source_commit_id = _source_commit(conn, asset["REPOSITORY_ID"])
        now = database._utcnow()
        conn.execute(
            """INSERT INTO EUC_ANALYSIS_RUNS
               (ANALYSIS_ID,EUC_ID,SOURCE_COMMIT_ID,ANALYZER_VERSION,STATUS,PROGRESS,CURRENT_STEP,STARTED_AT)
               VALUES (?,?,?,?,'VALIDATING',5,'FILE_VALIDATION',?)""",
            (analysis_id, euc_id, source_commit_id, settings.euc_analysis_version, now),
        )
        conn.execute("UPDATE EUC_ASSETS SET ANALYSIS_STATUS='VALIDATING',UPDATED_AT=? WHERE EUC_ID=?", (now, euc_id))
        conn.commit()
        record_audit_event("EUC_ANALYSIS_STARTED", actor_user_id=user_id,
                           repository_id=asset["REPOSITORY_ID"],
                           payload={"euc_id": euc_id, "analysis_id": analysis_id,
                                    "analyzer_version": settings.euc_analysis_version})
        payload = ledger_for_connection(conn).objects.get_bytes(conn, asset["ORIGINAL_OBJECT_HASH"])
        validate_euc_file(asset["ORIGINAL_FILENAME"], payload)
        suffix = f".{asset['FILE_TYPE']}"
        with NamedTemporaryFile(delete=False, suffix=suffix) as handle:
            handle.write(payload)
            temp_path = Path(handle.name)
        workbook = _workbook_for(temp_path, asset["FILE_TYPE"])
        if time.perf_counter() - started > settings.euc_analysis_timeout_seconds:
            raise EUCValidationError("ANALYSIS_TIMEOUT", "Workbook loading exceeded the analysis timeout.")
        context = AnalysisContext(euc_id, analysis_id, temp_path, asset["FILE_TYPE"],
                                  asset["REPOSITORY_ID"], _stable_sheets(conn, asset["REPOSITORY_ID"]), workbook)
        conn.execute("UPDATE EUC_ANALYSIS_RUNS SET STATUS='PARSING',PROGRESS=25,CURRENT_STEP='WORKBOOK_ANALYSIS' WHERE ANALYSIS_ID=?", (analysis_id,))
        conn.commit()
        workbook_result = WorkbookAnalyzer().analyze(context)
        if time.perf_counter() - started > settings.euc_analysis_timeout_seconds:
            raise EUCValidationError("ANALYSIS_TIMEOUT", "Workbook inventory exceeded the analysis timeout.")
        if workbook_result.metrics.get("used_cells", 0) > settings.euc_max_cells:
            raise EUCValidationError("OVERSIZED_WORKBOOK", "Workbook cell count exceeds the configured analysis limit.")
        formula_result = FormulaAnalyzer().analyze(context, workbook_result)
        if formula_result.metrics["total_formulas"] > settings.euc_max_formulas:
            raise EUCValidationError("OVERSIZED_WORKBOOK", "Workbook formula count exceeds the configured analysis limit.")
        openxml_result = OpenXMLAnalyzer().analyze(context)
        if time.perf_counter() - started > settings.euc_analysis_timeout_seconds:
            raise EUCValidationError("ANALYSIS_TIMEOUT", "Package inventory exceeded the analysis timeout.")
        workbook.close()
        results = [workbook_result, formula_result, openxml_result]
        warnings = [warning for result in results for warning in result.warnings]
        summary = _summary(asset, results)
        conn.execute("UPDATE EUC_ANALYSIS_RUNS SET STATUS='PERSISTING',PROGRESS=75,CURRENT_STEP='INVENTORY_PERSISTENCE' WHERE ANALYSIS_ID=?", (analysis_id,))
        _persist_inventory(conn, analysis_id, results, warnings)
        manifest = {"analysis_version": settings.euc_analysis_version, "euc_id": euc_id,
                    "source_commit_id": source_commit_id, "structure_hash": asset["STRUCTURE_HASH"],
                    "summary": summary, "sheets": workbook_result.records["sheets"],
                    "formula_patterns": formula_result.records["patterns"],
                    "external_links": openxml_result.records["external_links"],
                    "connections": openxml_result.records["connections"],
                    "objects": workbook_result.records["objects"] + openxml_result.records["objects"],
                    "warnings": [warning.__dict__ for warning in warnings]}
        manifest_object = ledger_for_connection(conn).objects.put(conn, "EUC_ANALYSIS_MANIFEST", manifest)
        conn.execute(
            "INSERT OR IGNORE INTO OBJECT_REFERENCES (PARENT_HASH,CHILD_HASH,REFERENCE_TYPE,POSITION) VALUES (?,?,?,0)",
            (manifest_object["object_hash"], asset["ORIGINAL_OBJECT_HASH"], "EUC_SOURCE"),
        )
        duration_ms = round((time.perf_counter() - started) * 1000)
        status = "COMPLETED_WITH_WARNINGS" if warnings else "COMPLETED"
        completed = database._utcnow()
        conn.execute(
            """UPDATE EUC_ANALYSIS_RUNS SET STATUS=?,PROGRESS=100,CURRENT_STEP='COMPLETED',
               COMPLETED_AT=?,DURATION_MS=?,WARNING_COUNT=?,RESULT_MANIFEST_HASH=?,SUMMARY_JSON=?,WARNINGS_JSON=?
               WHERE ANALYSIS_ID=?""",
            (status, completed, duration_ms, len(warnings), manifest_object["object_hash"],
             json.dumps(summary, sort_keys=True), json.dumps([warning.__dict__ for warning in warnings], sort_keys=True), analysis_id),
        )
        conn.execute("UPDATE EUC_ASSETS SET ANALYSIS_STATUS=?,LATEST_ANALYSIS_ID=?,UPDATED_AT=? WHERE EUC_ID=?",
                     (status, analysis_id, completed, euc_id))
        conn.commit()
        record_metric("euc_analysis_duration", duration_ms, "ms", repository_id=asset["REPOSITORY_ID"], status=status)
        record_audit_event("EUC_ANALYSIS_COMPLETED", actor_user_id=user_id,
                           repository_id=asset["REPOSITORY_ID"],
                           payload={"euc_id": euc_id, "analysis_id": analysis_id,
                                    "manifest_hash": manifest_object["object_hash"], "summary": summary}, status="SUCCESS")
        return get_analysis(euc_id, analysis_id, user_id)
    except Exception as exc:
        conn.rollback()
        try:
            completed = database._utcnow()
            conn.execute("UPDATE EUC_ANALYSIS_RUNS SET STATUS='FAILED',COMPLETED_AT=?,ERROR_COUNT=1,ERRORS_JSON=? WHERE ANALYSIS_ID=?",
                         (completed, json.dumps([{"code": getattr(exc, "code", "ANALYSIS_FAILURE"), "message": str(exc)}]), analysis_id))
            conn.execute("UPDATE EUC_ASSETS SET ANALYSIS_STATUS='FAILED',UPDATED_AT=? WHERE EUC_ID=?", (completed, euc_id))
            conn.commit()
        except Exception:
            conn.rollback()
        record_audit_event("EUC_ANALYSIS_FAILED", actor_user_id=user_id,
                           payload={"euc_id": euc_id, "analysis_id": analysis_id},
                           status="FAILED", failure_reason=str(exc))
        raise
    finally:
        if temp_path:
            temp_path.unlink(missing_ok=True)
        conn.close()


def _summary(asset, results) -> dict[str, Any]:
    workbook, formulas, package = results
    return {"filename": asset["ORIGINAL_FILENAME"], "file_type": asset["FILE_TYPE"],
            "size_bytes": asset["SIZE_BYTES"], **workbook.metrics, **formulas.metrics,
            **package.metrics, "macro_enabled": asset["FILE_TYPE"] == "xlsm" and package.metrics["vba_present"]}


def _persist_inventory(conn, analysis_id: str, results, warnings: list[AnalysisWarning]) -> None:
    workbook, formulas, package = results
    persisted_pattern_ids = {
        pattern["pattern_id"]: f"{analysis_id}_{pattern['normalized_hash'][:16].upper()}"
        for pattern in formulas.records["patterns"]
    }
    conn.executemany(
        """INSERT INTO EUC_SHEET_INVENTORY VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        [( _id("ESI"), analysis_id, s["sheet_id"], s["sheet_name"], s["position"], s["visibility"],
           s["max_row"], s["max_column"], s["used_range"], s["used_cells"], s["formula_cells"],
           s["constant_cells"], s["blank_styled_cells"], s["merged_ranges"], s["hidden_rows"],
           s["hidden_columns"], s["tables"], s["charts"], s["pivots"], s["validations"],
           s["comments"], s["hyperlinks"], s["structure_hash"]) for s in workbook.records["sheets"]],
    )
    conn.executemany(
        """INSERT INTO EUC_FORMULA_PATTERNS
           (PATTERN_ID,ANALYSIS_ID,NORMALIZED_FORMULA,NORMALIZED_HASH,OCCURRENCE_COUNT,
            FUNCTION_COUNT,FUNCTIONS_JSON,CROSS_SHEET_REFERENCE,EXTERNAL_REFERENCE,VOLATILE)
           VALUES (?,?,?,?,?,?,?,?,?,?)""",
        [(persisted_pattern_ids[p["pattern_id"]], analysis_id, p["normalized_formula"], p["normalized_hash"], p["occurrence_count"],
          p["function_count"], json.dumps(p["functions"]), int(p["cross_sheet"]), int(p["external"]), int(p["volatile"]))
         for p in formulas.records["patterns"]],
    )
    occurrences = formulas.records["occurrences"]
    for start in range(0, len(occurrences), 2000):
        conn.executemany(
            """INSERT INTO EUC_FORMULA_OCCURRENCES
               (OCCURRENCE_ID,ANALYSIS_ID,PATTERN_ID,SHEET_ID,CELL_ADDRESS,RAW_FORMULA)
               VALUES (?,?,?,?,?,?)""",
            [(_id("FOC"), analysis_id, persisted_pattern_ids[o["pattern_id"]], o["sheet_id"], o["cell_address"], o["raw_formula"])
             for o in occurrences[start:start + 2000]],
        )
    objects = workbook.records["objects"] + package.records["objects"]
    conn.executemany(
        """INSERT INTO EUC_OBJECT_INVENTORY
           (INVENTORY_ID,ANALYSIS_ID,SHEET_ID,OBJECT_TYPE,OBJECT_NAME,CELL_OR_RANGE,DETAILS_JSON,OBJECT_HASH)
           VALUES (?,?,?,?,?,?,?,?)""",
        [(_id("EOI"), analysis_id, item.get("sheet_id"), item["object_type"], item.get("name"), item.get("range"),
          json.dumps(item.get("details", {}), default=str, sort_keys=True), item.get("object_hash")) for item in objects],
    )
    conn.executemany(
        """INSERT INTO EUC_EXTERNAL_LINKS
           (LINK_ID,ANALYSIS_ID,SOURCE_EUC_REFERENCE,SOURCE_SHEET,SOURCE_ADDRESS,TARGET_SHEET_ID,TARGET_CELL,RELATIONSHIP_TYPE)
           VALUES (?,?,?,?,?,?,?,?)""",
        [(_id("ELK"), analysis_id, item.get("source_euc_reference"), item.get("source_sheet"), item.get("source_address"),
          item.get("target_sheet_id"), item.get("target_cell"), item["relationship_type"]) for item in package.records["external_links"]],
    )
    conn.executemany(
        """INSERT INTO EUC_CONNECTIONS
           (CONNECTION_ID,ANALYSIS_ID,CONNECTION_NAME,CONNECTION_TYPE,CREDENTIAL_PRESENT,DETAILS_JSON)
           VALUES (?,?,?,?,?,?)""",
        [(_id("ECN"), analysis_id, item.get("name"), item.get("type"), int(item.get("credential_present", False)),
          json.dumps(item.get("details", {}), sort_keys=True)) for item in package.records["connections"]],
    )
    conn.executemany(
        """INSERT INTO EUC_WARNINGS
           (WARNING_ID,ANALYSIS_ID,WARNING_CODE,SEVERITY,SHEET_ID,MESSAGE,DETAILS_JSON)
           VALUES (?,?,?,?,?,?,?)""",
        [(_id("EWN"), analysis_id, warning.code, warning.severity, warning.sheet_id,
          warning.message, json.dumps(warning.details, sort_keys=True)) for warning in warnings],
    )


def list_assets(
    user_id: str, repository_id: str | None = None, search: str = "",
    since: str | None = None, until: str | None = None,
) -> list[dict[str, Any]]:
    conn = database._get_connection()
    try:
        params: list[Any] = [user_id, user_id]
        filters = ["R.STATUS='ACTIVE'", "(R.CREATED_BY=? OR M.USER_ID=?)"]
        if repository_id:
            filters.append("A.REPOSITORY_ID=?")
            params.append(repository_id)
        if search:
            filters.append("(A.ORIGINAL_FILENAME LIKE ? OR R.REPOSITORY_NAME LIKE ?)")
            params.extend([f"%{search}%", f"%{search}%"])
        if since:
            filters.append("A.CREATED_AT>=?")
            params.append(since)
        if until:
            filters.append("A.CREATED_AT<=?")
            params.append(until)
        rows = conn.execute(
            f"""SELECT A.*,R.REPOSITORY_NAME,R.TABLE_ID,X.SUMMARY_JSON,X.ANALYZER_VERSION,X.WARNING_COUNT,X.COMPLETED_AT
                FROM EUC_ASSETS A JOIN WORKBOOK_REPOSITORIES R ON R.REPOSITORY_ID=A.REPOSITORY_ID
                LEFT JOIN REPOSITORY_MEMBERS M ON M.REPOSITORY_ID=R.REPOSITORY_ID AND M.USER_ID=?
                LEFT JOIN EUC_ANALYSIS_RUNS X ON X.ANALYSIS_ID=A.LATEST_ANALYSIS_ID
                WHERE {' AND '.join(filters)} ORDER BY A.UPDATED_AT DESC""",
            [user_id, *params],
        ).fetchall()
        return [{**_asset_dict(row), "repository_name": row["REPOSITORY_NAME"], "table_id": row["TABLE_ID"],
                 "summary": json.loads(row["SUMMARY_JSON"] or "{}"), "analyzer_version": row["ANALYZER_VERSION"],
                 "warning_count": row["WARNING_COUNT"] or 0, "completed_at": row["COMPLETED_AT"]} for row in rows]
    finally:
        conn.close()


def get_analysis(euc_id: str, analysis_id: str, user_id: str) -> dict[str, Any]:
    conn = database._get_connection()
    try:
        row = conn.execute(
            """SELECT X.*,A.REPOSITORY_ID,A.ORIGINAL_FILENAME,R.CREATED_BY,
                      B.HEAD_COMMIT_ID AS CURRENT_HEAD
               FROM EUC_ANALYSIS_RUNS X JOIN EUC_ASSETS A ON A.EUC_ID=X.EUC_ID
               JOIN WORKBOOK_REPOSITORIES R ON R.REPOSITORY_ID=A.REPOSITORY_ID
               LEFT JOIN BRANCHES B ON B.BRANCH_ID=R.DEFAULT_BRANCH_ID
               WHERE X.EUC_ID=? AND X.ANALYSIS_ID=?""",
            (euc_id, analysis_id),
        ).fetchone()
        if not row:
            raise KeyError("Analysis run does not exist.")
        _repository_access(conn, row["REPOSITORY_ID"], user_id)
        freshness = "CURRENT" if row["SOURCE_COMMIT_ID"] == row["CURRENT_HEAD"] else "STALE"
        if row["STATUS"] == "FAILED":
            freshness = "FAILED"
        elif row["STATUS"] == "COMPLETED_WITH_WARNINGS":
            freshness = "PARTIAL"
        return {"analysis_id": row["ANALYSIS_ID"], "euc_id": row["EUC_ID"], "status": row["STATUS"],
                "progress": row["PROGRESS"], "current_step": row["CURRENT_STEP"],
                "analyzer_version": row["ANALYZER_VERSION"], "source_commit_id": row["SOURCE_COMMIT_ID"],
                "current_head_commit_id": row["CURRENT_HEAD"], "freshness": freshness,
                "started_at": row["STARTED_AT"], "completed_at": row["COMPLETED_AT"],
                "duration_ms": row["DURATION_MS"], "warning_count": row["WARNING_COUNT"],
                "error_count": row["ERROR_COUNT"], "manifest_hash": row["RESULT_MANIFEST_HASH"],
                "summary": json.loads(row["SUMMARY_JSON"] or "{}"),
                "warnings": json.loads(row["WARNINGS_JSON"] or "[]"),
                "errors": json.loads(row["ERRORS_JSON"] or "[]")}
    finally:
        conn.close()


def get_inventory(euc_id: str, user_id: str, include: str, limit: int, cursor: int) -> dict[str, Any]:
    conn = database._get_connection()
    try:
        asset = conn.execute("SELECT * FROM EUC_ASSETS WHERE EUC_ID=?", (euc_id,)).fetchone()
        if not asset:
            raise KeyError("EUC asset does not exist.")
        _repository_access(conn, asset["REPOSITORY_ID"], user_id)
        analysis_id = asset["LATEST_ANALYSIS_ID"]
        if not analysis_id:
            return {"euc_id": euc_id, "status": asset["ANALYSIS_STATUS"], "items": [], "next_cursor": None}
        run = conn.execute("SELECT * FROM EUC_ANALYSIS_RUNS WHERE ANALYSIS_ID=?", (analysis_id,)).fetchone()
        sections = {item.strip() for item in include.split(",") if item.strip()} or {"overview"}
        response: dict[str, Any] = {"euc_id": euc_id, "analysis_id": analysis_id,
                                   "status": run["STATUS"], "overview": json.loads(run["SUMMARY_JSON"] or "{}")}
        if "sheets" in sections:
            response["sheets"] = [{key.lower(): row[key] for key in row.keys()} for row in conn.execute(
                "SELECT * FROM EUC_SHEET_INVENTORY WHERE ANALYSIS_ID=? ORDER BY SHEET_POSITION LIMIT ? OFFSET ?",
                (analysis_id, limit, cursor))]
        if "formulas" in sections:
            response["formulas"] = [{**{key.lower(): row[key] for key in row.keys()},
                                      "functions": json.loads(row["FUNCTIONS_JSON"])} for row in conn.execute(
                "SELECT * FROM EUC_FORMULA_PATTERNS WHERE ANALYSIS_ID=? ORDER BY OCCURRENCE_COUNT DESC LIMIT ? OFFSET ?",
                (analysis_id, limit, cursor))]
        if "objects" in sections:
            response["objects"] = [{**{key.lower(): row[key] for key in row.keys()},
                                     "details": json.loads(row["DETAILS_JSON"])} for row in conn.execute(
                "SELECT * FROM EUC_OBJECT_INVENTORY WHERE ANALYSIS_ID=? ORDER BY OBJECT_TYPE,OBJECT_NAME LIMIT ? OFFSET ?",
                (analysis_id, limit, cursor))]
        if "dependencies" in sections:
            response["external_links"] = [{key.lower(): row[key] for key in row.keys()} for row in conn.execute(
                "SELECT * FROM EUC_EXTERNAL_LINKS WHERE ANALYSIS_ID=? LIMIT ? OFFSET ?", (analysis_id, limit, cursor))]
            response["connections"] = [{**{key.lower(): row[key] for key in row.keys()},
                                         "details": json.loads(row["DETAILS_JSON"])} for row in conn.execute(
                "SELECT * FROM EUC_CONNECTIONS WHERE ANALYSIS_ID=? LIMIT ? OFFSET ?", (analysis_id, limit, cursor))]
        response["next_cursor"] = cursor + limit if any(len(response.get(key, [])) == limit for key in ("sheets", "formulas", "objects", "external_links")) else None
        return response
    finally:
        conn.close()


def resolve_external_link(euc_id: str, user_id: str, link_id: str, resolution_status: str) -> dict[str, Any]:
    """Mark an EUC external-link dependency RESOLVED (or back to UNRESOLVED)
    once someone has actually repointed or accounted for it. Previously
    there was no write path at all for RESOLUTION_STATUS -- every external
    link stayed 'UNRESOLVED' forever, so migration readiness scoring
    (``migration.readiness.blockers``) permanently treated it as a CRITICAL
    blocker even after the dependency was actually addressed."""
    resolution_status = resolution_status.upper()
    if resolution_status not in {"RESOLVED", "UNRESOLVED"}:
        raise ValueError("resolution_status must be RESOLVED or UNRESOLVED")
    conn = database._get_connection()
    try:
        asset = conn.execute("SELECT REPOSITORY_ID FROM EUC_ASSETS WHERE EUC_ID=?", (euc_id,)).fetchone()
        if not asset:
            raise KeyError("EUC asset does not exist")
        _repository_access(conn, asset["REPOSITORY_ID"], user_id, edit=True)
        link = conn.execute(
            """SELECT L.LINK_ID FROM EUC_EXTERNAL_LINKS L JOIN EUC_ANALYSIS_RUNS A ON A.ANALYSIS_ID=L.ANALYSIS_ID
               WHERE L.LINK_ID=? AND A.EUC_ID=?""",
            (link_id, euc_id),
        ).fetchone()
        if not link:
            raise KeyError("External link does not exist for this asset")
        conn.execute("UPDATE EUC_EXTERNAL_LINKS SET RESOLUTION_STATUS=? WHERE LINK_ID=?", (resolution_status, link_id))
        conn.commit()
    finally:
        conn.close()
    record_audit_event(
        "EUC_EXTERNAL_LINK_RESOLUTION_CHANGED", actor_user_id=user_id,
        payload={"euc_id": euc_id, "link_id": link_id, "resolution_status": resolution_status},
    )
    return {"link_id": link_id, "resolution_status": resolution_status}


def export_manifest(euc_id: str, user_id: str) -> tuple[str, bytes]:
    conn = database._get_connection()
    try:
        asset = conn.execute("SELECT * FROM EUC_ASSETS WHERE EUC_ID=?", (euc_id,)).fetchone()
        if not asset or not asset["LATEST_ANALYSIS_ID"]:
            raise KeyError("A completed analysis is required before export.")
        _repository_access(conn, asset["REPOSITORY_ID"], user_id)
        run = conn.execute("SELECT RESULT_MANIFEST_HASH FROM EUC_ANALYSIS_RUNS WHERE ANALYSIS_ID=?", (asset["LATEST_ANALYSIS_ID"],)).fetchone()
        manifest = ledger_for_connection(conn).objects.get(conn, run[0])
        return f"{Path(asset['ORIGINAL_FILENAME']).stem}_inventory.json", json.dumps(manifest, indent=2, sort_keys=True).encode()
    finally:
        conn.close()
