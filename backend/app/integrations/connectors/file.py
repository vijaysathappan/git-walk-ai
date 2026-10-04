"""Safe CSV, JSON, XML, and Excel connector backed by immutable source objects."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import defusedxml.ElementTree as DET
from datetime import datetime, timezone
from typing import Any

from openpyxl import load_workbook

from .base import Connector, ConnectorError, ConnectorReadResult, ConnectorRecord


class FileConnector(Connector):
    connector_type = "FILE"
    capabilities = frozenset({"READ", "WRITE", "DISCOVER_SCHEMA", "BULK_WRITE"})

    def _payload(self) -> bytes:
        object_hash = self.configuration.get("object_hash")
        if not object_hash or not self.object_loader:
            raise ConnectorError("INVALID_CONFIGURATION", "Upload a source file before running this connection")
        return self.object_loader(object_hash)

    def _records(self) -> list[dict[str, Any]]:
        payload = self._payload()
        file_format = str(self.configuration.get("format", "CSV")).upper()
        encoding = self.configuration.get("encoding", "utf-8-sig")
        if file_format == "CSV":
            reader = csv.DictReader(io.StringIO(payload.decode(encoding)), delimiter=self.configuration.get("delimiter", ","))
            return [dict(row) for row in reader]
        if file_format == "JSON":
            value = json.loads(payload.decode(encoding))
            records = value if isinstance(value, list) else value.get(self.configuration.get("records_path", "records"), [value])
            return [dict(item) for item in records]
        if file_format == "XML":
            root = DET.fromstring(payload)
            row_tag = self.configuration.get("row_tag")
            nodes = root.findall(f".//{row_tag}") if row_tag else list(root)
            return [{child.tag: child.text for child in node} for node in nodes]
        if file_format == "EXCEL":
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            try:
                sheet = workbook[self.configuration.get("sheet_name")] if self.configuration.get("sheet_name") else workbook.active
                rows = sheet.iter_rows(values_only=True)
                headers = [str(value or f"COLUMN_{index + 1}") for index, value in enumerate(next(rows, ())) ]
                return [dict(zip(headers, row)) for row in rows]
            finally:
                workbook.close()
        raise ConnectorError("INVALID_CONFIGURATION", f"Unsupported file format {file_format}")

    async def test_connection(self) -> dict[str, Any]:
        payload = self._payload()
        return {"status": "SUCCESS", "bytes": len(payload), "format": self.configuration.get("format", "CSV").upper(), "capabilities": sorted(self.capabilities)}

    async def discover(self) -> dict[str, Any]:
        records = self._records()
        fields: dict[str, dict[str, Any]] = {}
        for record in records[:100]:
            for key, value in record.items():
                inferred = "NULL" if value is None else type(value).__name__.upper()
                current = fields.setdefault(key, {"name": key, "types": set(), "nullable": False})
                current["types"].add(inferred)
                current["nullable"] = current["nullable"] or value is None or value == ""
        return {"object_type": self.configuration.get("object_type", "Record"), "record_count_sampled": min(100, len(records)),
                "fields": [{**item, "types": sorted(item["types"])} for item in fields.values()]}

    async def read(self, request: dict[str, Any]) -> ConnectorReadResult:
        raw = self._payload(); records = self._records(); start = int(request.get("checkpoint", {}).get("offset", 0)); batch_size = int(request.get("batch_size", 1000))
        selected = records[start:start + batch_size]; id_field = request.get("id_field") or self.configuration.get("id_field")
        version_field = request.get("version_field") or self.configuration.get("version_field")
        output = []
        for index, payload in enumerate(selected, start=start):
            external_id = str(payload.get(id_field)) if id_field and payload.get(id_field) is not None else str(index + 1)
            canonical = json.dumps(payload, sort_keys=True, default=str, separators=(",", ":"))
            version = str(payload.get(version_field)) if version_field and payload.get(version_field) is not None else hashlib.sha256(canonical.encode()).hexdigest()[:20]
            output.append(ConnectorRecord(self.configuration.get("object_type", "Record"), external_id, version, payload,
                                          metadata={"source_row": index + 2, "format": self.configuration.get("format", "CSV").upper()}))
        return ConnectorReadResult(output, {"offset": start + len(selected), "complete": start + len(selected) >= len(records)}, len(raw))

    async def write(self, request: dict[str, Any]) -> dict[str, Any]:
        records = request.get("records") or []
        file_format = str(self.configuration.get("format", "JSON")).upper()
        if file_format == "CSV":
            output = io.StringIO(); fields = sorted({key for row in records for key in row})
            writer = csv.DictWriter(output, fieldnames=fields); writer.writeheader(); writer.writerows(records)
            payload = output.getvalue().encode(self.configuration.get("encoding", "utf-8"))
        else:
            payload = json.dumps(records, default=str, separators=(",", ":")).encode("utf-8")
        return {"status": "COMPLETED", "records_written": len(records), "payload": payload,
                "content_type": "text/csv" if file_format == "CSV" else "application/json"}
